"""Pro-forma credit notes, written into the forwarding team's own Excel template.

static/templates/credit_note_template.xlsx is their "Credit Note List - Forwarding"
workbook with every filled-in value cleared: a "List" sheet (the CREDIT NOTE REQUEST
with up to 12 lines) and one "CN" sheet (the PRO-FORMA CREDIT NOTE). For a request
the List is filled in and the CN sheet is copied once per credit note, so the result
looks exactly like the workbook they fill in by hand.
"""
import copy
import io
import zipfile
from datetime import date
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.drawing.image import Image

TEMPLATE = Path(__file__).parent / "static" / "templates" / "credit_note_template.xlsx"
LIST_ROWS = 12  # request lines on one List sheet (rows 4–15)
CREDIT_CODE = 30050
CREDIT_DESCRIPTION = "FORWARDING"


def _num(v):
    """File and invoice numbers go in as numbers, like on the hand-made sheets."""
    s = str(v or "").strip()
    return int(s) if s.isdigit() else s


def _logo_bytes() -> bytes:
    with zipfile.ZipFile(TEMPLATE) as z:
        name = next(n for n in z.namelist() if n.startswith("xl/media/"))
        return z.read(name)


def _copy_logo(src, dst, data: bytes) -> None:
    for img in src._images:
        new = Image(io.BytesIO(data))
        new.anchor = copy.deepcopy(img.anchor)
        new.width, new.height = img.width, img.height
        dst.add_image(new)


def _copy_print_setup(src, dst) -> None:
    dst.page_setup.paperSize = src.page_setup.paperSize
    dst.page_setup.orientation = src.page_setup.orientation
    dst.page_setup.fitToWidth = src.page_setup.fitToWidth
    dst.page_setup.fitToHeight = src.page_setup.fitToHeight
    dst.page_setup.scale = src.page_setup.scale
    dst.sheet_properties.pageSetUpPr = copy.deepcopy(src.sheet_properties.pageSetUpPr)
    dst.page_margins = copy.deepcopy(src.page_margins)
    dst.print_options = copy.deepcopy(src.print_options)
    dst.sheet_view.showGridLines = src.sheet_view.showGridLines


# Rows that grow to fit their text on the original sheets (openpyxl would otherwise
# save them at a fixed height, squashing e.g. "(Please print name)")
AUTO_ROWS = {"List": (2,), "CN": (29, 31)}


def _auto_rows(ws, rows) -> None:
    for r in rows:
        ws.row_dimensions[r].height = None


def _fill_note(ws, day: date, n: dict) -> None:
    amount = round(float(n.get("amount") or 0), 2)
    # "Updated On" and its date aren't shown on the credit note
    ws["K1"] = None
    ws["K2"] = None
    # Room for the two-line labels ("File Nº to be / credited:", "Invoice Nº to be / credited:")
    # and the "Information That the Accounts Department…" heading, in any font
    for r, h in ((4, 42), (5, 42), (8, 38)):
        ws.row_dimensions[r].height = h
    ws["B3"] = f"Date:         {day.strftime('%d/%m/%Y')}"
    ws["D5"] = _num(n.get("file"))
    ws["J5"] = _num(n.get("fs_invoice"))
    ws["D6"] = f"FORWARDING RECOVERED ON INVOICE {str(n.get('client_invoice') or '').strip()}".rstrip()
    ws["B11"] = CREDIT_CODE
    ws["C11"] = CREDIT_DESCRIPTION
    ws["I11"] = amount
    ws["J11"] = 0
    ws["I25"] = amount
    ws["J25"] = 0
    ws["I26"] = amount
    if amount >= 100000:  # R100,000.00 and up no longer fits the Amount column: borrow from VAT (R0.00)
        ws.column_dimensions["I"].width = 17
        ws.column_dimensions["J"].width = 15.1


def _fill_list(ws, day: date, notes: list[dict]) -> None:
    ws["B1"] = day.strftime("%Y/%m/%d")
    for i, n in enumerate(notes):
        r = 4 + i
        ws[f"B{r}"] = round(float(n.get("amount") or 0), 2)
        ws[f"C{r}"] = _num(n.get("fs_invoice"))
        ws[f"D{r}"] = _num(n.get("file"))
        ws[f"E{r}"] = "Requested"
    ws["C16"] = round(sum(float(n.get("amount") or 0) for n in notes), 2)
    ws["E26"] = f"Date:         {day.strftime('%Y/%m/%d')}"


def build_credit_workbook(day: date, notes: list[dict]) -> bytes:
    """notes: [{file, fs_invoice, client_invoice, amount}] → the request workbook."""
    wb = load_workbook(TEMPLATE, rich_text=True)  # keeps the raised "st" in "Authorised 1st Director"
    list_tpl, cn_tpl = wb["List"], wb["CN"]

    # One List per 12 credit notes (the template has 12 request lines)
    chunks = [notes[i:i + LIST_ROWS] for i in range(0, len(notes), LIST_ROWS)] or [[]]
    lists = [list_tpl]
    for k in range(1, len(chunks)):
        extra = wb.copy_worksheet(list_tpl)
        extra.title = f"List ({k + 1})"
        _copy_print_setup(list_tpl, extra)
        lists.append(extra)
    for ws, chunk in zip(lists, chunks):
        _fill_list(ws, day, chunk)

    # Copy the blank CN sheet first, then fill each one in
    logo = _logo_bytes() if len(notes) > 1 else b""
    sheets = [cn_tpl] if notes else []
    for _ in notes[1:]:
        ws = wb.copy_worksheet(cn_tpl)
        _copy_logo(cn_tpl, ws, logo)
        _copy_print_setup(cn_tpl, ws)
        sheets.append(ws)
    for ws, n in zip(sheets, notes):
        _fill_note(ws, day, n)
        _auto_rows(ws, AUTO_ROWS["CN"])
    for ws in lists:
        _auto_rows(ws, AUTO_ROWS["List"])
    if not notes:
        wb.remove(cn_tpl)
    for i, ws in enumerate(sheets):
        ws.title = f"CN-{i + 1}"

    wb._sheets = lists + sheets  # List first, then CN-1, CN-2, …
    wb.active = 0
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
