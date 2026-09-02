"""Iteration 6 tests: append-comment endpoint, date prefix refresh, XLSX layout."""
import io
import os
from datetime import date

import pytest
import requests
from openpyxl import load_workbook
from openpyxl.drawing.spreadsheet_drawing import OneCellAnchor

def _load_backend_url() -> str:
    v = os.environ.get("REACT_APP_BACKEND_URL")
    if not v:
        env = os.path.join(os.path.dirname(__file__), "..", "..", "frontend", ".env")
        with open(env) as f:
            for line in f:
                if line.startswith("REACT_APP_BACKEND_URL="):
                    v = line.split("=", 1)[1].strip()
                    break
    return v.rstrip("/")


BASE_URL = _load_backend_url()
API = f"{BASE_URL}/api"

FLANGE_ID = "61aa6ac3-3adb-44a7-832b-94799a99faf7"  # Patuma
CF_ID = "b2d8c8cc-f981-4ea8-954c-cef5422cf3ce"  # Clearfreight


def today_prefix() -> str:
    t = date.today()
    return f"{t.day:02d}.{t.month:02d}"


@pytest.fixture(scope="module")
def created_shipments():
    ids = []
    yield ids
    for sid in ids:
        try:
            requests.delete(f"{API}/shipments/{sid}", timeout=10)
        except Exception:
            pass


def _create_shipment(client_id: str, planned_etd: str = "2025-09-12") -> dict:
    r = requests.post(f"{API}/shipments", json={
        "client_id": client_id,
        "supplier": "TEST_iter6",
        "planned_etd": planned_etd,
    }, timeout=15)
    assert r.status_code == 200, r.text
    return r.json()


def test_create_shipment_seeds_today_prefix(created_shipments):
    s = _create_shipment(FLANGE_ID)
    created_shipments.append(s["id"])
    prefix = today_prefix()
    assert s["comments"].startswith(f"{prefix} - Planned ETD "), s["comments"]
    assert "12.09" in s["comments"]


def test_append_comment_preserves_today_prefix(created_shipments):
    s = _create_shipment(FLANGE_ID)
    created_shipments.append(s["id"])
    r = requests.post(f"{API}/shipments/{s['id']}/append-comment",
                      json={"snippet": "Awaiting confirmation of departure."}, timeout=15)
    assert r.status_code == 200, r.text
    doc = r.json()
    prefix = today_prefix()
    assert doc["comments"].startswith(f"{prefix} - "), doc["comments"]
    assert "Planned ETD" in doc["comments"]
    assert "Awaiting confirmation of departure." in doc["comments"]
    # Separator ". "
    assert ". Awaiting confirmation of departure." in doc["comments"]


def test_append_comment_idempotent(created_shipments):
    s = _create_shipment(FLANGE_ID)
    created_shipments.append(s["id"])
    snippet = "Awaiting confirmation of departure."
    r1 = requests.post(f"{API}/shipments/{s['id']}/append-comment", json={"snippet": snippet}, timeout=15)
    r2 = requests.post(f"{API}/shipments/{s['id']}/append-comment", json={"snippet": snippet}, timeout=15)
    assert r1.status_code == 200 and r2.status_code == 200
    assert r2.json()["comments"].count(snippet) == 1


def test_append_comment_404():
    r = requests.post(f"{API}/shipments/nonexistent-id-xyz/append-comment",
                      json={"snippet": "hi"}, timeout=10)
    assert r.status_code == 404


def test_mark_delayed_refreshes_prefix_and_appends(created_shipments):
    s = _create_shipment(FLANGE_ID)
    created_shipments.append(s["id"])
    # Simulate an older prefix by directly PATCHing an old-date-prefixed comment
    requests.patch(f"{API}/shipments/{s['id']}", json={
        "comments": "01.01 - Planned ETD 12.09."
    }, timeout=10)
    r = requests.post(f"{API}/shipments/{s['id']}/mark-delayed",
                      json={"new_etd": "2025-09-18"}, timeout=15)
    assert r.status_code == 200, r.text
    c = r.json()["comments"]
    prefix = today_prefix()
    assert c.startswith(f"{prefix} - "), c
    assert not c.startswith("01.01"), c
    assert "Vessel delayed slightly. Now planned ETD 18.09" in c


def test_mark_shipped_refreshes_prefix_and_appends(created_shipments):
    s = _create_shipment(FLANGE_ID)
    created_shipments.append(s["id"])
    requests.patch(f"{API}/shipments/{s['id']}", json={
        "comments": "01.01 - Planned ETD 12.09."
    }, timeout=10)
    r = requests.post(f"{API}/shipments/{s['id']}/mark-shipped",
                      json={"sob_date": "2025-09-18"}, timeout=15)
    assert r.status_code == 200, r.text
    c = r.json()["comments"]
    prefix = today_prefix()
    assert c.startswith(f"{prefix} - "), c
    assert "SOB 18.09" in c and "Awaiting ANF." in c


def test_full_comment_pipeline(created_shipments):
    """End-to-end: create -> append -> delayed -> shipped."""
    s = _create_shipment(FLANGE_ID)
    created_shipments.append(s["id"])
    requests.post(f"{API}/shipments/{s['id']}/append-comment",
                  json={"snippet": "Awaiting confirmation of departure."}, timeout=15)
    requests.post(f"{API}/shipments/{s['id']}/mark-delayed",
                  json={"new_etd": "2025-09-18"}, timeout=15)
    r = requests.post(f"{API}/shipments/{s['id']}/mark-shipped",
                      json={"sob_date": "2025-09-18"}, timeout=15)
    c = r.json()["comments"]
    prefix = today_prefix()
    assert c.startswith(f"{prefix} - ")
    assert "Planned ETD 12.09" in c
    assert "Awaiting confirmation of departure." in c
    assert "Vessel delayed slightly. Now planned ETD 18.09" in c
    assert "SOB 18.09" in c
    assert "Awaiting ANF." in c


def _load_xlsx(client_id: str):
    r = requests.get(f"{API}/reports/{client_id}/xlsx", timeout=30)
    assert r.status_code == 200, r.text
    return load_workbook(io.BytesIO(r.content))


def test_xlsx_flange_layout():
    wb = _load_xlsx(FLANGE_ID)
    ws = wb.active

    # Row 1 height = 80
    assert ws.row_dimensions[1].height == 80

    # Row 4 header bg 0033CC and font 33CCFF
    hcell = ws.cell(row=4, column=1)
    assert (hcell.fill.fgColor.rgb or "").endswith("0033CC"), hcell.fill.fgColor.rgb
    assert (hcell.font.color.rgb or "").endswith("33CCFF"), hcell.font.color.rgb

    # Image embedded
    assert len(ws._images) >= 1, "no image embedded"
    img = ws._images[0]
    assert isinstance(img.anchor, OneCellAnchor), f"anchor type: {type(img.anchor)}"

    # Column widths sum <= 200
    widths = []
    for col_letter in [c.column_letter for c in ws[4] if c.value]:
        w = ws.column_dimensions[col_letter].width
        if w:
            widths.append(w)
    total = sum(widths)
    assert total <= 200, f"total column width {total} > 200"

    # Comments column width == 32; find it by header label
    comments_col_idx = None
    for cell in ws[4]:
        if cell.value == "Comments":
            comments_col_idx = cell.column_letter
            break
    assert comments_col_idx is not None
    assert ws.column_dimensions[comments_col_idx].width == 32


def test_xlsx_clearfreight_centered_logo():
    wb = _load_xlsx(CF_ID)
    ws = wb.active
    assert len(ws._images) >= 1
    assert isinstance(ws._images[0].anchor, OneCellAnchor)
    assert ws.row_dimensions[1].height == 80
