"""Iteration 2 backend tests: auto-delayed comments, dashboard reminders, report styling."""
import os
from datetime import date, timedelta

import pytest
import requests
from dotenv import load_dotenv
from pathlib import Path

load_dotenv(Path(__file__).parents[2] / "frontend" / ".env")
BASE_URL = os.environ["REACT_APP_BACKEND_URL"].rstrip("/")


@pytest.fixture(scope="module")
def api():
    s = requests.Session()
    s.headers.update({"Content-Type": "application/json"})
    return s


@pytest.fixture(scope="module")
def client_id(api):
    r = api.post(f"{BASE_URL}/api/clients", json={"name": "TEST_Iter2_Client"})
    assert r.status_code == 200, r.text
    cid = r.json()["id"]
    yield cid
    api.delete(f"{BASE_URL}/api/clients/{cid}")


def _mkship(api, client_id, **kw):
    payload = {"client_id": client_id, "supplier": "TEST", "comments": "Initial booking.", "status": "Booked"}
    payload.update(kw)
    r = api.post(f"{BASE_URL}/api/shipments", json=payload)
    assert r.status_code == 200, r.text
    return r.json()


# ---- Auto-delayed comment tests ----
def test_etd_delay_appends_comment_and_sets_delayed(api, client_id):
    today = date.today()
    etd_initial = (today + timedelta(days=2)).isoformat()
    ship = _mkship(api, client_id, planned_etd=etd_initial, comments="Initial booking.")
    new_etd = (today + timedelta(days=6)).isoformat()
    r = api.patch(f"{BASE_URL}/api/shipments/{ship['id']}", json={"planned_etd": new_etd})
    assert r.status_code == 200, r.text
    data = r.json()
    assert data["status"] == "Delayed"
    assert "Vessel delayed slightly. Now planned ETD" in data["comments"]
    d = date.fromisoformat(new_etd)
    assert f"{d.day:02d}.{d.month:02d}." in data["comments"]


def test_vessel_name_change_appends_and_delays(api, client_id):
    ship = _mkship(api, client_id, vessel_name="MSC Alpha", comments="ok.")
    r = api.patch(f"{BASE_URL}/api/shipments/{ship['id']}", json={"vessel_name": "MSC Beta"})
    assert r.status_code == 200
    data = r.json()
    assert data["status"] == "Delayed"
    assert "Vessel changed by S/Line to MSC Beta" in data["comments"]


def test_first_time_etd_set_does_not_trigger_delay(api, client_id):
    ship = _mkship(api, client_id, planned_etd=None, comments="Fresh.")
    future = (date.today() + timedelta(days=5)).isoformat()
    r = api.patch(f"{BASE_URL}/api/shipments/{ship['id']}", json={"planned_etd": future})
    assert r.status_code == 200
    data = r.json()
    assert "Vessel delayed slightly" not in data["comments"]
    assert data["status"] == "Booked"


# ---- Reminders tests ----
def test_reminders_lcl_10d(api, client_id):
    eta = (date.today() + timedelta(days=8)).isoformat()
    ship = _mkship(api, client_id, cargo_type="LCL", planned_eta=eta, supplier="TEST_LCL")
    r = api.get(f"{BASE_URL}/api/dashboard/reminders")
    assert r.status_code == 200
    rems = r.json()["reminders"]
    match = [x for x in rems if x["shipment_id"] == ship["id"]]
    assert match, f"Expected LCL reminder for {ship['id']}"
    assert match[0]["cargo_type"] == "LCL"
    assert "LCL cargo report" in match[0]["kind"]


def test_reminders_fcl_2d(api, client_id):
    etd = (date.today() + timedelta(days=1)).isoformat()
    ship = _mkship(api, client_id, cargo_type="FCL", planned_etd=etd, supplier="TEST_FCL")
    r = api.get(f"{BASE_URL}/api/dashboard/reminders")
    match = [x for x in r.json()["reminders"] if x["shipment_id"] == ship["id"]]
    assert match
    assert "before departure" in match[0]["kind"]


def test_reminders_fcl_second_vessel(api, client_id):
    etd1 = (date.today() + timedelta(days=30)).isoformat()  # far
    etd2 = (date.today() + timedelta(days=1)).isoformat()
    ship = _mkship(api, client_id, cargo_type="FCL", planned_etd=etd1,
                   second_vessel_name="MSC Second", second_vessel_etd=etd2, supplier="TEST_2ND")
    r = api.get(f"{BASE_URL}/api/dashboard/reminders")
    match = [x for x in r.json()["reminders"] if x["shipment_id"] == ship["id"]]
    assert match
    assert "2nd vessel" in match[0]["kind"]
    assert match[0]["target_date"] == etd2


def test_reminders_excludes_anf(api, client_id):
    etd = (date.today() + timedelta(days=1)).isoformat()
    ship = _mkship(api, client_id, cargo_type="FCL", planned_etd=etd, supplier="TEST_ANF")
    api.patch(f"{BASE_URL}/api/shipments/{ship['id']}", json={"anf_received": True})
    r = api.get(f"{BASE_URL}/api/dashboard/reminders")
    match = [x for x in r.json()["reminders"] if x["shipment_id"] == ship["id"]]
    assert not match


# ---- Reports ----
def test_xlsx_and_pdf_magic(api, client_id):
    r = api.get(f"{BASE_URL}/api/reports/{client_id}/xlsx")
    assert r.status_code == 200
    assert "spreadsheetml" in r.headers["content-type"]
    assert r.content[:2] == b"PK"
    r = api.get(f"{BASE_URL}/api/reports/{client_id}/pdf")
    assert r.status_code == 200
    assert "application/pdf" in r.headers["content-type"]
    assert r.content[:5] == b"%PDF-"


def test_report_preview_vessel_block_second_vessel(api, client_id):
    ship = _mkship(api, client_id,
                   vessel_name="MSC Alpha",
                   second_vessel_name="MSC Beta",
                   tracking_doc_number="MEDU999",
                   carrier="MSC",
                   supplier="TEST_VBLK")
    r = api.get(f"{BASE_URL}/api/reports/{client_id}/preview")
    assert r.status_code == 200
    data = r.json()
    row = next(row for row in data["rows"] if row.get("supplier") == "TEST_VBLK")
    vb = row["vessel_block"]
    assert "MSC Alpha (1st Vessel)" in vb
    assert "MSC Beta (2nd Vessel)" in vb
    assert "MEDU999" in vb
    assert "[MSC]" in vb


def test_xlsx_header_styling(api, client_id):
    """Verify header uses light gray D9D9D9 fill and black text/borders."""
    import io
    from openpyxl import load_workbook
    r = api.get(f"{BASE_URL}/api/reports/{client_id}/xlsx")
    wb = load_workbook(io.BytesIO(r.content))
    ws = wb.active
    header_cell = ws.cell(row=2, column=1)
    # Fill color
    fg = header_cell.fill.fgColor.rgb or ""
    assert "D9D9D9" in fg, f"Header fill was {fg}, expected D9D9D9"
    # Text color black
    fc = header_cell.font.color.rgb or ""
    assert "000000" in fc, f"Header font color was {fc}"
    # Border black
    assert header_cell.border.left.color.rgb.endswith("000000")
