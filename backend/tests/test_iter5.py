"""Iteration 5 tests: logos in PDF/XLSX, vessel_block no carrier, Planned status, comments column widths."""
import io
import os
import zipfile

import openpyxl
import pytest
import requests
from dotenv import load_dotenv
load_dotenv("/app/frontend/.env")

BASE_URL = os.environ["REACT_APP_BACKEND_URL"].rstrip("/")
API = f"{BASE_URL}/api"


@pytest.fixture(scope="module")
def clients():
    r = requests.get(f"{API}/clients", timeout=30)
    r.raise_for_status()
    return r.json()


@pytest.fixture(scope="module")
def patuma_client(clients):
    for c in clients:
        if c.get("company") == "Patuma":
            return c
    pytest.skip("No Patuma client")


@pytest.fixture(scope="module")
def cf_client(clients):
    for c in clients:
        if c.get("company") == "Clearfreight":
            return c
    pytest.skip("No Clearfreight client")


# ---- Companies API ----
def test_companies_have_logos():
    r = requests.get(f"{API}/companies", timeout=15)
    assert r.status_code == 200
    companies = {c["key"]: c for c in r.json()["companies"]}
    assert "Patuma" in companies and "Clearfreight" in companies
    for key in ("Patuma", "Clearfreight"):
        logo = companies[key].get("logo")
        assert logo, f"{key} missing logo"
        assert os.path.exists(logo), f"logo file missing: {logo}"


# ---- Vessel block no carrier tag ----
def test_vessel_block_no_carrier_bracket(patuma_client):
    # Create a shipment
    payload = {
        "client_id": patuma_client["id"],
        "supplier": "TEST_iter5",
        "vessel_name": "Marianna",
        "tracking_doc_number": "MEDUADK75346",
        "carrier": "MSC",
    }
    r = requests.post(f"{API}/shipments", json=payload, timeout=15)
    assert r.status_code == 200
    ship_id = r.json()["id"]
    try:
        preview = requests.get(f"{API}/reports/{patuma_client['id']}/preview", timeout=15).json()
        target_row = None
        for row in preview["rows"]:
            if "Marianna" in (row.get("vessel_block") or ""):
                target_row = row
                break
        assert target_row is not None
        vb = target_row["vessel_block"]
        assert "[MSC]" not in vb
        assert "[" not in vb
        assert vb == "Marianna\nMEDUADK75346"
    finally:
        requests.delete(f"{API}/shipments/{ship_id}", timeout=15)


# ---- Planned status ----
def test_planned_status_roundtrip(patuma_client):
    payload = {"client_id": patuma_client["id"], "supplier": "TEST_planned", "status": "Planned"}
    r = requests.post(f"{API}/shipments", json=payload, timeout=15)
    assert r.status_code == 200
    ship = r.json()
    assert ship["status"] == "Planned"
    ship_id = ship["id"]
    try:
        # patch to Planned as well
        r2 = requests.patch(f"{API}/shipments/{ship_id}", json={"status": "Planned"}, timeout=15)
        assert r2.status_code == 200
        assert r2.json()["status"] == "Planned"
    finally:
        requests.delete(f"{API}/shipments/{ship_id}", timeout=15)


# ---- PDF has embedded logo (large size) ----
def test_pdf_embeds_logo_patuma(patuma_client):
    r = requests.get(f"{API}/reports/{patuma_client['id']}/pdf", timeout=30)
    assert r.status_code == 200
    assert r.headers["content-type"].startswith("application/pdf")
    size = len(r.content)
    print(f"Patuma PDF size: {size} bytes")
    assert size > 50_000, f"PDF too small ({size}), likely missing embedded logo"


def test_pdf_embeds_logo_clearfreight(cf_client):
    r = requests.get(f"{API}/reports/{cf_client['id']}/pdf", timeout=30)
    assert r.status_code == 200
    size = len(r.content)
    print(f"Clearfreight PDF size: {size} bytes")
    assert size > 40_000


# ---- XLSX embeds logo image ----
def test_xlsx_embeds_logo(patuma_client):
    r = requests.get(f"{API}/reports/{patuma_client['id']}/xlsx", timeout=30)
    assert r.status_code == 200
    assert "spreadsheetml" in r.headers["content-type"]
    size = len(r.content)
    print(f"XLSX size: {size} bytes")
    assert size > 50_000
    # Verify /xl/media/ contains an image
    zf = zipfile.ZipFile(io.BytesIO(r.content))
    media_files = [n for n in zf.namelist() if n.startswith("xl/media/")]
    print(f"Media files: {media_files}")
    assert any(n.endswith((".png", ".jpg", ".jpeg")) for n in media_files), "No embedded image in xlsx"


# ---- Excel Comments column widths + wrap + row heights ----
def test_xlsx_comments_column_layout(patuma_client):
    long_comment = "Awaiting confirmation of departure. H/bill released by supplier. Extra text here to force wrap."
    payload = {
        "client_id": patuma_client["id"],
        "supplier": "TEST_widecomment",
        "vessel_name": "TestVsl",
        "tracking_doc_number": "TESTDOC123",
        "carrier": "MSC",
        "comments": long_comment,
    }
    r = requests.post(f"{API}/shipments", json=payload, timeout=15)
    assert r.status_code == 200
    ship_id = r.json()["id"]
    try:
        r2 = requests.get(f"{API}/reports/{patuma_client['id']}/xlsx", timeout=30)
        assert r2.status_code == 200
        wb = openpyxl.load_workbook(io.BytesIO(r2.content))
        ws = wb.active

        # Find Comments column at row 4 (header row)
        comments_col = None
        header_map = {}
        for col_idx in range(1, ws.max_column + 1):
            v = ws.cell(row=4, column=col_idx).value
            if v:
                header_map[v] = col_idx
            if v == "Comments":
                comments_col = col_idx
        assert comments_col is not None, f"Comments column not found: headers={header_map}"

        # (a) Comments column width >= 42
        from openpyxl.utils import get_column_letter
        col_letter = get_column_letter(comments_col)
        width = ws.column_dimensions[col_letter].width
        print(f"Comments column ({col_letter}) width: {width}")
        assert width >= 42, f"Comments column width {width} < 42"

        # Find our data row with TEST_widecomment
        data_row = None
        supplier_col = header_map.get("Supplier")
        for row_idx in range(5, ws.max_row + 1):
            if ws.cell(row=row_idx, column=supplier_col).value == "TEST_widecomment":
                data_row = row_idx
                break
        assert data_row is not None

        # (b) row height >= 48
        rh = ws.row_dimensions[data_row].height
        print(f"Row {data_row} height: {rh}")
        assert rh is not None and rh >= 48

        # (c) wrap_text True on comments cell
        cmt_cell = ws.cell(row=data_row, column=comments_col)
        assert cmt_cell.alignment.wrap_text is True
        assert long_comment in (cmt_cell.value or "")

        # Neighboring columns separate: Copy Docs Status / H/bill values not merged into comments
        # Just verify comments cell value equals the long comment (no bleed of other cells)
        assert cmt_cell.value == long_comment
    finally:
        requests.delete(f"{API}/shipments/{ship_id}", timeout=15)
