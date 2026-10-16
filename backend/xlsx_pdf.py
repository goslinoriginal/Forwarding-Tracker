"""Draws the credit note sheets straight to PDF, the way Excel prints them.

Render has no Excel or LibreOffice, so this reads the sheet's own layout — column
widths, row heights, merged cells, borders, fonts (incl. superscripts like the "st" in
"1st"), number formats, alignment, the logo, page breaks and fit-to-width — and draws
it with ReportLab. It covers what the forwarding team's credit note workbook uses;
it isn't a general spreadsheet renderer.
"""
import colorsys
import io
import re
import xml.etree.ElementTree as ET
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.utils import get_column_letter, range_boundaries
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.utils import ImageReader
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas

FONT_DIR = Path(__file__).parent / "static" / "fonts"
EMU_PER_PX = 9525
MDW = 8  # max digit width (px) of the workbook's default font, Arial 11
PAD = 2.0  # pt, Excel's inner cell margin

_registered = False


def _register_fonts() -> None:
    global _registered
    if _registered:
        return
    for fam, base in (("LS", "LiberationSans"), ("CA", "Carlito")):
        for style, suffix in (("", "Regular"), ("B", "Bold"), ("I", "Italic"), ("BI", "BoldItalic")):
            pdfmetrics.registerFont(TTFont(f"{fam}{style}", str(FONT_DIR / f"{base}-{suffix}.ttf")))
    _registered = True


def _font_name(name: str, bold: bool, italic: bool) -> str:
    fam = "LS" if (name or "").lower() in ("arial", "helvetica", "liberation sans") else "CA"
    return fam + ("B" if bold else "") + ("I" if italic else "")


# ---------- colours ----------
def _theme_palette(wb) -> list[str]:
    """Excel's theme colour order: lt1, dk1, lt2, dk2, accent1..6, hlink, folHlink."""
    default = ["FFFFFF", "000000", "E7E6E6", "44546A", "4472C4", "ED7D31", "A5A5A5", "FFC000", "5B9BD5", "70AD47", "0563C1", "954F72"]
    try:
        root = ET.fromstring(wb.loaded_theme)
    except Exception:
        return default
    ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
    scheme = root.find(".//a:clrScheme", ns)
    if scheme is None:
        return default
    vals = {}
    for el in scheme:
        tag = el.tag.split("}")[1]
        child = el[0]
        vals[tag] = child.get("lastClr") or child.get("val") or "000000"
    order = ["lt1", "dk1", "lt2", "dk2", "accent1", "accent2", "accent3", "accent4", "accent5", "accent6", "hlink", "folHlink"]
    # Theme XML lists dk1 before lt1; Excel indexes lt1 first
    return [vals.get(k, d) for k, d in zip(order, default)]


def _tint(hex6: str, tint: float) -> str:
    r, g, b = (int(hex6[i:i + 2], 16) / 255 for i in (0, 2, 4))
    h, l, s = colorsys.rgb_to_hls(r, g, b)
    l = l * (1 + tint) if tint < 0 else l * (1 - tint) + tint
    r, g, b = colorsys.hls_to_rgb(h, max(0, min(1, l)), s)
    return "".join(f"{round(v * 255):02X}" for v in (r, g, b))


def _color(c, palette, default="000000") -> str:
    try:
        if c is None:
            return default
        if c.type == "rgb" and isinstance(c.rgb, str):
            return c.rgb[-6:]
        if c.type == "theme" and c.theme is not None:
            return _tint(palette[c.theme], c.tint or 0) if c.theme < len(palette) else default
    except Exception:
        pass
    return default


def _rgb(hex6: str):
    return tuple(int(hex6[i:i + 2], 16) / 255 for i in (0, 2, 4))


# ---------- values ----------
def _format(value, fmt: str) -> str:
    if value is None:
        return ""
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        m = re.match(r'"([^"]*)"#,##0(\.0+)?', fmt or "")
        if m:
            dec = len(m.group(2) or ".")-1
            return f"{m.group(1)}{value:,.{dec}f}"
        if isinstance(value, float) and not value.is_integer():
            return f"{value:g}"
        return str(int(value))
    if hasattr(value, "strftime"):
        return value.strftime("%Y/%m/%d") if "yyyy" in (fmt or "") else value.strftime("%m-%d-%y")
    return str(value)


def _runs(cell, palette):
    """[(text, font, size, colour, superscript, underline)] for plain or rich text."""
    f = cell.font
    base = (_font_name(f.name, f.b, f.i), float(f.sz or 11), _color(f.color, palette), False, bool(f.u))
    v = cell.value
    if isinstance(v, CellRichText):
        out = []
        for part in v:
            if isinstance(part, TextBlock):
                fi = part.font
                name = fi.rFont or f.name
                out.append((part.text, _font_name(name, fi.b if fi.b is not None else f.b, fi.i if fi.i is not None else f.i),
                            float(fi.sz or f.sz or 11), _color(fi.color, palette, base[2]), fi.vertAlign == "superscript", bool(fi.u) or base[4]))
            else:
                out.append((str(part), *base))
        return out
    text = _format(v, cell.number_format)
    return [(text, *base)] if text else []


def _wrap(runs, width):
    """Splits runs into lines no wider than `width` (pt), breaking between words."""
    lines, line, lw = [], [], 0.0
    for text, font, size, col, sup, und in runs:
        sz = size * (0.62 if sup else 1)
        for i, para in enumerate(text.split("\n")):
            if i:
                lines.append(line); line, lw = [], 0.0
            wrapped = False
            for word in re.split(r"(\s+)", para):
                if not word:
                    continue
                w = pdfmetrics.stringWidth(word, font, sz)
                if lw + w > width and line and not word.isspace():
                    lines.append(line); line, lw = [], 0.0
                    wrapped = True
                if not line and word.isspace() and wrapped:
                    continue  # spaces at a line break are dropped; leading ones are kept
                line.append((word, font, sz, col, sup, und, w)); lw += w
    lines.append(line)
    return lines


# ---------- drawing ----------
_WIDTH = {"hair": 0.25, "thin": 0.5, "medium": 1.0, "thick": 1.6, "double": 0.5, "dotted": 0.5, "dashed": 0.5, "mediumDashed": 1.0}


def _line(c, x1, y1, x2, y2, side, palette):
    if not side or not side.style:
        return
    c.setStrokeColorRGB(*_rgb(_color(side.color, palette)))
    c.setLineWidth(_WIDTH.get(side.style, 0.5))
    c.setDash([1, 1] if side.style == "dotted" else [3, 2] if "ashed" in side.style else [])
    if side.style == "double":
        dx, dy = (0, 1.0) if y1 == y2 else (1.0, 0)
        c.line(x1 - dx, y1 - dy, x2 - dx, y2 - dy)
        c.line(x1 + dx, y1 + dy, x2 + dx, y2 + dy)
    else:
        c.line(x1, y1, x2, y2)
    c.setDash([])


def _print_range(ws):
    if ws.print_area:
        ref = ws.print_area if isinstance(ws.print_area, str) else ws.print_area[0]
        ref = ref.split("!")[-1].replace("$", "")
        return range_boundaries(ref)
    return range_boundaries(ws.dimensions)


def draw_sheet(c, ws, palette, page_w, page_h, last_col=None):
    min_col, min_row, max_col, max_row = _print_range(ws)
    # A column setting can cover several columns (e.g. B:C); the page runs to the last sized column,
    # like Excel's print of text that spills past the used cells
    dims = {}
    for d in ws.column_dimensions.values():
        for ci in range(d.min or 0, (d.max or d.min or 0) + 1):
            dims[ci] = d
    if not ws.print_area:
        # Like the approved prints: from A1, one narrow spacer column past the used cells
        min_col, min_row = 1, 1
        nxt = dims.get(max_col + 1)
        if nxt is not None and nxt.width and nxt.width < 4:
            max_col += 1
    if last_col:
        max_col = min(max_col, last_col)
    default_w = ws.sheet_format.defaultColWidth or (ws.sheet_format.baseColWidth or 8) + 0.43
    col_w = {}
    for ci in range(min_col, max_col + 1):
        d = dims.get(ci)
        w = d.width if d is not None and d.width else default_w
        col_w[ci] = 0 if (d is not None and d.hidden) else (w * MDW + 5) * 0.75  # px → pt
    default_h = ws.sheet_format.defaultRowHeight or 15
    row_h = {}
    for r in range(min_row, max_row + 1):
        rd = ws.row_dimensions[r]
        row_h[r] = 0 if rd.hidden else (rd.height if rd.height is not None else default_h)

    m = ws.page_margins
    left, right, top, bottom = m.left * 72, m.right * 72, m.top * 72, m.bottom * 72
    content_w = sum(col_w.values())
    fit = bool(ws.sheet_properties.pageSetUpPr and ws.sheet_properties.pageSetUpPr.fitToPage)
    avail = page_w - left - right
    # "Fit to page width" fills the width (as the approved prints did); otherwise only shrink
    scale = (avail / content_w if fit else min(1.0, avail / content_w)) if content_w else 1.0
    x0 = left + max(0.0, (page_w - left - right - content_w * scale) / 2) if ws.print_options.horizontalCentered else left

    # Pages: split at manual row breaks
    breaks = sorted(b.id for b in ws.row_breaks.brk if min_row <= b.id < max_row)
    pages, start = [], min_row
    for b in breaks:
        pages.append((start, b)); start = b + 1
    pages.append((start, max_row))

    merged = {}
    covered = set()
    in_merge = {}  # (row, col) → the merged range it belongs to
    for mr in ws.merged_cells.ranges:
        merged[(mr.min_row, mr.min_col)] = (mr.max_row, mr.max_col)
        for r in range(mr.min_row, mr.max_row + 1):
            for cc in range(mr.min_col, mr.max_col + 1):
                in_merge[(r, cc)] = mr
                if (r, cc) != (mr.min_row, mr.min_col):
                    covered.add((r, cc))

    for pi, (r0, r1) in enumerate(pages):
        if pi:
            c.showPage()
        c.saveState()
        c.translate(x0, page_h - top)
        c.scale(scale, scale)
        # cell geometry (y grows downward from the page's first row)
        xs, x = {}, 0.0
        for ci in range(min_col, max_col + 2):
            xs[ci] = x; x += col_w.get(ci, 0)
        ys, y = {}, 0.0
        for r in range(r0, r1 + 2):
            ys[r] = y; y += row_h.get(r, 0)

        def box(r, ci):
            rr, cc = merged.get((r, ci), (r, ci))
            rr, cc = min(rr, r1), min(cc, max_col)
            return xs[ci], ys[r], xs[cc + 1] - xs[ci], ys[rr + 1] - ys[r]

        # borders
        for r in range(r0, r1 + 1):
            for ci in range(min_col, max_col + 1):
                cell = ws.cell(row=r, column=ci)
                if not cell.has_style:
                    continue
                bx, by, bw, bh = xs[ci], ys[r], col_w[ci], row_h[r]
                if not bw or not bh:
                    continue
                b = cell.border
                mr = in_merge.get((r, ci))  # inside a merged cell only its outer edge is drawn
                if not mr or r == mr.min_row:
                    _line(c, bx, -by, bx + bw, -by, b.top, palette)
                if not mr or r == mr.max_row:
                    _line(c, bx, -(by + bh), bx + bw, -(by + bh), b.bottom, palette)
                if not mr or ci == mr.min_col:
                    _line(c, bx, -by, bx, -(by + bh), b.left, palette)
                if not mr or ci == mr.max_col:
                    _line(c, bx + bw, -by, bx + bw, -(by + bh), b.right, palette)

        # text
        for r in range(r0, r1 + 1):
            for ci in range(min_col, max_col + 1):
                if (r, ci) in covered:
                    continue
                cell = ws.cell(row=r, column=ci)
                runs = _runs(cell, palette)
                if not runs:
                    continue
                bx, by, bw, bh = box(r, ci)
                if not bh:
                    continue  # hidden row
                al = cell.alignment
                is_num = isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool)
                h = al.horizontal or ("right" if is_num else "left")
                lines = _wrap(runs, bw - 2 * PAD) if al.wrap_text else [[(t, f, s * (0.62 if sp else 1), co, sp, u, pdfmetrics.stringWidth(t, f, s * (0.62 if sp else 1))) for t, f, s, co, sp, u in runs]]
                big = max(s for _, _, s, *_ in runs)
                lh = big * 1.17
                total = lh * len(lines)
                v = al.vertical or "bottom"
                ty = by + (PAD if v == "top" else (bh - total) / 2 if v in ("center", "justify", "distributed") else bh - total - PAD * 0.5)
                for li, line in enumerate(lines):
                    lw = sum(p[6] for p in line)
                    tx = bx + (bw - lw) / 2 if h in ("center", "centerContinuous") else bx + bw - PAD - lw if h == "right" else bx + PAD
                    base = -(ty + lh * (li + 1) - big * 0.22)
                    last = max((i for i, p in enumerate(line) if not p[0].isspace()), default=-1)
                    for i, (word, font, sz, col, sup, und, w) in enumerate(line):
                        c.setFillColorRGB(*_rgb(col))
                        c.setFont(font, sz)
                        yy = base + (big * 0.36 if sup else 0)
                        c.drawString(tx, yy, word)
                        if und and i <= last:
                            c.setStrokeColorRGB(*_rgb(col)); c.setLineWidth(max(0.5, sz / 16))
                            c.line(tx, base - big * 0.12, tx + w, base - big * 0.12)
                        tx += w

        # pictures (the logo)
        for img in ws._images:
            a = img.anchor
            fr = a._from
            if not (r0 - 1 <= fr.row <= r1 - 1):
                continue
            ix = xs.get(fr.col + 1, 0) + fr.colOff / EMU_PER_PX * 0.75
            iy = ys.get(fr.row + 1, 0) + fr.rowOff / EMU_PER_PX * 0.75
            to = getattr(a, "to", None)
            if to is not None:
                iw = xs.get(to.col + 1, 0) + to.colOff / EMU_PER_PX * 0.75 - ix
                ih = ys.get(to.row + 1, 0) + to.rowOff / EMU_PER_PX * 0.75 - iy
            else:
                iw, ih = a.ext.width / EMU_PER_PX * 0.75, a.ext.height / EMU_PER_PX * 0.75
            data = img._data() if not hasattr(img, "_cached") else img._cached
            c.drawImage(ImageReader(io.BytesIO(data)), ix, -(iy + ih), iw, ih, mask="auto")
        c.restoreState()


def workbook_pdf(xlsx: bytes, last_col: dict | None = None) -> bytes:
    """Every sheet of the workbook, in order, as one PDF (A4 like the sheets)."""
    _register_fonts()
    wb = load_workbook(io.BytesIO(xlsx), rich_text=True)
    palette = _theme_palette(wb)
    out = io.BytesIO()
    c = canvas.Canvas(out, pagesize=A4)
    for i, ws in enumerate(wb.worksheets):
        size = landscape(A4) if ws.page_setup.orientation == "landscape" else A4
        c.setPageSize(size)
        if i:
            c.showPage()
            c.setPageSize(size)
        draw_sheet(c, ws, palette, *size, last_col=(last_col or {}).get(ws.title))
    c.save()
    return out.getvalue()
