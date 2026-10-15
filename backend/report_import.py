"""Read the old Excel status reports (Patuma / Clearfreight) into clients + shipments.

Every report is one client: a date and "Client: <name>" above a header row starting
with "Supplier", then one row per shipment. Columns are matched by header text, so
reports with extra columns (SOB date, copy docs, freight rate...) work the same way.

Dates in the comments are dd.mm without a year; they're placed in the year that
puts them closest to the report date.
"""
from __future__ import annotations

import io
import re
from datetime import date, datetime
from typing import Optional

from openpyxl import load_workbook

# header text (lower-case, punctuation removed) -> field
HEADER_ALIASES = [
    ("supplier", "supplier"),
    ("order", "order_booking_file"),
    ("booking file", "file_number"),
    ("file", "file_number"),
    ("shipped pending", "status"),
    ("status", "status"),
    ("sob date", "sob_date"),
    ("sailing date", "sob_date"),
    ("shipped on board", "sob_date"),
    ("sob", "sob_date"),
    ("vessel", "vessel_block"),
    ("pol", "pol"),
    ("port of loading", "pol"),
    ("eta", "eta"),
    ("final destination", "final_destination"),
    ("destination", "final_destination"),
    ("comments", "comments"),
    ("copy docs", "copy_docs_status"),
    ("h bill", "hbill_released"),
    ("hbill", "hbill_released"),
    ("freight rate", "expected_freight_rate"),
    ("o freight", "expected_freight_rate"),
    ("rate", "expected_freight_rate"),
]
OPTIONAL_FIELDS = ["sob_date", "pol", "final_destination", "hbill_released", "copy_docs_status", "expected_freight_rate"]

DOC_CARRIER = [
    (re.compile(r"^MEDU"), "MSC"), (re.compile(r"^MSCU|^MSMU|^MSDU"), "MSC"),
    (re.compile(r"^MAEU|^MRKU|^MSKU"), "Maersk"),
    (re.compile(r"^ONEY|^ONEU"), "ONE"),
    (re.compile(r"^COSU|^CBHU|^CSNU"), "COSCO"),
    (re.compile(r"^HLCU|^HLXU"), "Hapag Lloyd"),
    (re.compile(r"^CMDU|^CMAU|^[A-Z]{3}\d{7}$"), "CMA CGM"),  # incl. bookings like CAD0911040
    (re.compile(r"^PCIU|^PILU"), "PIL"),
    # Consolidator house bills, e.g. TXGDUR6199781V / NBGJNB6723503V
    (re.compile(r"^[A-Z]{3}(DUR|JNB|CPT|PLZ|ELS)\d+V$"), "Vanguard"),
]
VESSEL_CARRIER = [("MSC", "MSC"), ("MAERSK", "Maersk"), ("CMA CGM", "CMA CGM"), ("COSCO", "COSCO"), ("ONE ", "ONE"), ("HAPAG", "Hapag Lloyd")]

DM = r"(\d{1,2})\.(\d{1,2})\.?"
LEG = r"(?:\s*\(\s*(\d)\s*(?:st|nd|rd|th)?\s*vessel\s*\))?"
SOB_RE = re.compile(r"\bSOB\s+" + DM + LEG, re.I)
ETD_RE = re.compile(r"\b(?:(?:now|provisionally)\s+)?pl?anned\s+ETD\s+" + DM + LEG, re.I)
PREFIX_RE = re.compile(r"^\s*" + DM + r"\s*-\s*")
FULL_DATE_RE = re.compile(r"^\s*(\d{1,2})[./](\d{1,2})[./](\d{4})\s*$")
DOC_RE = re.compile(r"^[A-Z0-9]{8,}$")
ARCHIVE_PATTERNS = [
    (re.compile(r"\bANF\s+rec(?:ei|ie)ved|\bANF\s+docs\s+to\s+ops", re.I), "ANF received"),
    (re.compile(r"\bSOB'?s\s+rec(?:ei|ie)ved|\bSOB'?s?\s+docs\s+to\s+ops", re.I), "SOB's received"),
]
# Sentences the system writes itself; anything else is kept as a note.
AUTO_SENTENCE = re.compile(
    r"^\(?\s*\d\s*(st|nd|rd|th)?\s*vessel\s*\)?\.?$|^SOB\b(?!'?s)|planned\s+ETD|delayed|vessel changed|ETD\s+TBA|"
    r"^Awaiting (ANF|SOB'?s|confirmation of departure)|ANF\s+rec|SOB'?s\s+rec|Docs to Ops|^p?l?anned$|\bETA\s+\d",
    re.I,
)


def _norm(v) -> str:
    return " ".join(str(v or "").replace("’", "'").split())


def _key(v) -> str:
    return re.sub(r"[^a-z ]", " ", _norm(v).lower()).strip()


def _field_for(header: str) -> Optional[str]:
    h = " ".join(_key(header).split())
    if not h:
        return None
    words = set(h.split())
    for alias, field in HEADER_ALIASES:
        if set(alias.split()) <= words:
            return field
    return None


def _cell_date(v) -> Optional[date]:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    m = FULL_DATE_RE.match(_norm(v))
    if m:
        try:
            return date(int(m.group(3)), int(m.group(2)), int(m.group(1)))
        except ValueError:
            return None
    return None


def _dm(day: str, month: str, ref: date) -> Optional[date]:
    """dd.mm -> the date in whichever year is closest to the report date."""
    best = None
    for y in (ref.year - 1, ref.year, ref.year + 1):
        try:
            d = date(y, int(month), int(day))
        except ValueError:
            continue
        if best is None or abs((d - ref).days) < abs((best - ref).days):
            best = d
    return best


def _carrier(doc: str, vessel: str) -> str:
    for rx, name in DOC_CARRIER:
        if doc and rx.search(doc):
            return name
    v = (vessel or "").upper() + " "
    for prefix, name in VESSEL_CARRIER:
        if v.startswith(prefix):
            return name
    return "Other"


HOUSE_BILL_RE = re.compile(r"^[A-Z]{3}(DUR|JNB|CPT|PLZ|ELS)\d+V$")


def _split_vessels(text: str) -> tuple[list[str], str, list[str]]:
    """'Wan Hai (1st Vessel), MSC Loretta (2nd Vessel) PKGJNB6170968V' -> (names, doc, other docs).
    When a house bill and the carrier's B/L are both listed, the carrier's B/L is the
    tracking number if its prefix identifies the carrier; otherwise the house bill."""
    t = _norm(text)
    docs: list[str] = []
    while True:
        tokens = [x for x in re.split(r"[\s,]+", t) if x]
        if tokens and DOC_RE.match(tokens[-1]) and re.search(r"\d", tokens[-1]):
            docs.insert(0, tokens[-1])
            t = t[: t.rfind(tokens[-1])].strip(" ,")
        else:
            break
    master = [d for d in docs if not HOUSE_BILL_RE.match(d) and any(rx.search(d) for rx, _ in DOC_CARRIER)]
    doc = master[0] if master else (docs[0] if docs else "")
    names = []
    for part in re.split(r",|\)\s+", t):
        name = re.sub(r"\(?\s*\d\s*(st|nd|rd|th)?\s*vessel\s*\)?", "", part, flags=re.I).strip(" ,()")
        if name:
            names.append(name)
    return names, doc, [d for d in docs if d != doc]


def _cargo_type(order: str, comments: str, docs: list[str]) -> str:
    o = order.upper()
    if re.search(r"\d\s*X\s*(20|40)|FCL", o):
        return "FCL"
    if "LCL" in o or re.search(r"SOB'?S", comments.upper()) or any(HOUSE_BILL_RE.match(d) for d in docs):
        return "LCL"
    if "CBM" in o and not re.search(r"(20|40)\s*'", o):
        return "LCL"
    return "FCL"


def parse_report(data: bytes, filename: str, cutoff: date) -> dict:
    wb = load_workbook(io.BytesIO(data), data_only=True)
    ws = wb.active
    warnings: list[str] = []

    rows = [[c.value for c in row] for row in ws.iter_rows()]
    header_idx = next((i for i, r in enumerate(rows) if any(_key(v) == "supplier" for v in r)), None)
    if header_idx is None:
        raise ValueError("No header row with 'Supplier' found")

    # Client name, report date and title above the header
    client_name, report_date, title = "", None, ""
    for r in rows[:header_idx]:
        for i, v in enumerate(r):
            k = _key(v)
            if k.startswith("client") and not client_name:
                client_name = next((_norm(x) for x in r[i + 1:] if _norm(x)), "")
            if report_date is None and _cell_date(v):
                report_date = _cell_date(v)
            if "report" in k and not title:
                title = _norm(v)
    if not client_name:
        client_name = re.sub(r"^[0-9a-f]{8}-", "", filename.rsplit(".", 1)[0]).replace("_", " ")
        warnings.append("Client name not found in the sheet — used the file name")
    if report_date is None:
        report_date = date.today()
        warnings.append("Report date not found — assumed today")

    # Columns: header row, falling back to the row above for headers merged across both
    header = rows[header_idx]
    above = rows[header_idx - 1] if header_idx > 0 else []
    cols: dict[int, str] = {}
    extra_cols: dict[int, str] = {}
    for i, v in enumerate(header):
        f = _field_for(v) or (_field_for(above[i]) if i < len(above) else None)
        if f and f not in cols.values():
            cols[i] = f
        elif _norm(v) and not f:
            extra_cols[i] = _norm(v)  # a column the system has no field for

    # Entity: a company name written in the heading wins; otherwise the Patuma
    # letterhead logo is much wider than Clearfreight's
    heading = " ".join(_norm(v).upper() for r in rows[:header_idx] for v in r if v)
    company = "Clearfreight"
    for img in ws._images:
        try:
            w, h = img.width, img.height
            if h and w / h > 8.5:
                company = "Patuma"
        except Exception:
            pass
    if "SEAFREIGHT FORWARDING" not in title.upper() and title.upper().strip() == "SHIPPING REPORT":
        company = "Patuma"
    if "PATUMA" in heading:
        company = "Patuma"
    elif "CLEARFREIGHT" in heading or "CLEAR FREIGHT" in heading:
        company = "Clearfreight"

    shipments = []
    for r in rows[header_idx + 1:]:
        rec = {f: _norm(r[i]) if i < len(r) else "" for i, f in cols.items()}
        extras = [f"{h}: {_norm(r[i])}" for i, h in extra_cols.items() if i < len(r) and _norm(r[i])]
        if extras:
            rec["order_booking_file"] = "\n".join([x for x in [rec.get("order_booking_file", "")] if x] + extras)
        if not (rec.get("supplier") or rec.get("file_number") or rec.get("order_booking_file")):
            continue
        shipments.append(_parse_row(rec, r, cols, report_date, cutoff))

    return {
        "file": filename,
        "client": {
            "name": client_name,
            "company": company,
            "optional_columns": {k: k in cols.values() for k in OPTIONAL_FIELDS},
        },
        "report_date": report_date.isoformat(),
        "shipments": shipments,
        "warnings": warnings,
    }


def _parse_row(rec: dict, raw: list, cols: dict, ref: date, cutoff: date) -> dict:
    comments = rec.get("comments", "")
    body = comments
    last_update = ref
    m = PREFIX_RE.match(comments)
    if m:
        last_update = _dm(m.group(1), m.group(2), ref) or ref
        body = comments[m.end():]

    names, doc, other_docs = _split_vessels(rec.get("vessel_block", ""))
    # LCL is tracked with Vanguard on the house bill, so prefer it over the carrier's B/L
    cargo = _cargo_type(rec.get("order_booking_file", ""), comments, [doc, *other_docs])
    if cargo == "LCL":
        house = next((d for d in [doc, *other_docs] if HOUSE_BILL_RE.match(d)), None)
        if house and house != doc:
            other_docs = [d for d in [doc, *other_docs] if d != house]
            doc = house
    legs = [{"name": n, "planned_etd": None, "sob_date": None, "state": "planned"} for n in names] or [{"name": "", "planned_etd": None, "sob_date": None, "state": "planned"}]

    def leg_at(n: int):
        while len(legs) < n:
            legs.append({"name": "", "planned_etd": None, "sob_date": None, "state": "planned"})
        return legs[n - 1]

    for sm in SOB_RE.finditer(body):
        d = _dm(sm.group(1), sm.group(2), ref)
        if d:
            leg_at(int(sm.group(3) or 1))["sob_date"] = d.isoformat()
    for em in ETD_RE.finditer(body):
        d = _dm(em.group(1), em.group(2), ref)
        if not d:
            continue
        n = int(em.group(3)) if em.group(3) else next((i + 1 for i, l in enumerate(legs) if not l["sob_date"]), 1)
        leg = leg_at(n)
        leg["planned_etd"] = d.isoformat()
        before = body[: em.start()][-80:].lower()
        if "delayed" in before:
            leg["state"] = "delayed"
        elif "changed" in before:
            leg["state"] = "changed"

    # SOB date column wins over the comment for the 1st vessel
    sob_col = rec.get("sob_date", "")
    if sob_col:
        idx = next(i for i, f in cols.items() if f == "sob_date")
        d = _cell_date(raw[idx])
        legs[0]["sob_date"] = d.isoformat() if d else sob_col

    eta = ""
    if rec.get("eta"):
        idx = next(i for i, f in cols.items() if f == "eta")
        d = _cell_date(raw[idx])
        eta = d.isoformat() if d else ""

    if not eta:
        em2 = re.search(r"\bETA\s+" + DM, body, re.I)
        if em2:
            d = _dm(em2.group(1), em2.group(2), ref)
            eta = d.isoformat() if d else ""

    hbill = None
    if re.search(r"H/?bill\s+NOT\s+released", comments, re.I):
        hbill = False
    elif re.search(r"H/?bill\s+released", comments, re.I):
        hbill = True
    hb_col = rec.get("hbill_released", "").lower()
    if hb_col in ("yes", "y"):
        hbill = True
    elif hb_col in ("no", "n"):
        hbill = False

    reason = next((label for rx, label in ARCHIVE_PATTERNS if rx.search(comments)), None)
    if not reason and "cancel" in rec.get("status", "").lower():
        reason = "Cancelled"
    if not reason and last_update < cutoff:
        reason = f"No update since {last_update.strftime('%d.%m')}"

    # Keep the sentences the system won't write itself (e.g. "H/bill released by supplier.")
    rest = ETD_RE.sub(" ", SOB_RE.sub(" ", body))
    rest = re.sub(r"\(?\s*\d\s*(st|nd|rd|th)?\s*vessel\s*\)?\.?", " ", rest, flags=re.I)
    sentences = [s.strip(" .") + "." for s in re.split(r"(?<=\.)\s+", rest) if s.strip(" .")]
    note = " ".join(" ".join(s.split()) for s in sentences if not AUTO_SENTENCE.search(s))

    order = rec.get("order_booking_file", "")
    second = legs[1] if len(legs) > 1 else None
    ship = {
        "supplier": rec.get("supplier", ""),
        "order_booking_file": "\n".join(line.strip(" ,") for line in order.split("\n")),
        "file_number": rec.get("file_number", ""),
        "cargo_type": cargo,
        "carrier": _carrier(doc, legs[0]["name"]),
        "vessel_name": legs[0]["name"],
        "planned_etd": legs[0]["planned_etd"],
        "etd_tba": bool(re.search(r"ETD\s+TBA", body, re.I)),
        "sob_date": legs[0]["sob_date"],
        "first_vessel_state": legs[0]["state"],
        "second_vessel_name": second["name"] if second else None,
        "second_vessel_etd": second["planned_etd"] if second else None,
        "second_vessel_sob_date": second["sob_date"] if second else None,
        "second_vessel_state": second["state"] if second else "planned",
        "extra_vessels": [{"name": l["name"], "planned_etd": l["planned_etd"], "sob_date": l["sob_date"], "state": l["state"]} for l in legs[2:]],
        "tracking_doc_number": doc,
        "eta": eta or None,
        "pol": rec.get("pol") or None,
        "final_destination": rec.get("final_destination") or None,
        "copy_docs_status": rec.get("copy_docs_status") or None,
        "expected_freight_rate": rec.get("expected_freight_rate") or None,
        "hbill_released": hbill,
        "comment_note": note,
        "original_comment": comments,
        "other_docs": other_docs,
        "last_update": last_update.isoformat(),
        "archived": bool(reason),
        "archive_reason": reason,
    }
    return ship
